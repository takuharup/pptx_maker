"""Tests for Phase 2: SQLite + JSON image mapping → VBA dataset."""
import json
import os
import sqlite3
import tempfile

import pytest
from pptx import Presentation
from pptx.util import Inches

from core.phase2_replace import DatasetBuilder, ImageMappingValidator, Phase2Pipeline
from core.validator import ValidationError

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

REQUIRED_SHAPES = [
    "TXT_ConfigFile",
    "TXT_Conditions",
    "TXT_Dispersion",
    "TXT_GasHU",
    "IMG_Main",
    "IMG_Graph1",
    "IMG_Graph2",
]


def _make_db(db_path: str) -> None:
    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    con = sqlite3.connect(db_path)
    con.execute(
        """CREATE TABLE IF NOT EXISTS cases (
            case_id TEXT PRIMARY KEY,
            config_file TEXT,
            conditions TEXT,
            max_height REAL,
            dispersion_height REAL,
            gas_holdup REAL,
            created_at TEXT
        )"""
    )
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?)",
        ("1a", "input_202604_case1a", "flat底面\nz=40 mm", 0.125, 0.088, 0.718, "2026-04-23 00:00:00"),
    )
    con.execute(
        "INSERT INTO cases VALUES (?,?,?,?,?,?,?)",
        ("1b", "input_202604_case1b", "flat底面\nz=60 mm", 0.138, 0.096, 0.742, "2026-04-23 00:00:00"),
    )
    con.commit()
    con.close()


def _make_pptx(pptx_path: str, shapes: list[str]) -> None:
    os.makedirs(os.path.dirname(pptx_path), exist_ok=True)
    prs = Presentation()
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)
    for name in shapes:
        txBox = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(0.3))
        txBox.name = name
        txBox.text_frame.text = f"{{{{{name}}}}}"
    prs.save(pptx_path)


def _make_mapping_json(path: str, img_dir: str) -> dict:
    """Create a minimal image_mapping.json with real (empty) image files."""
    os.makedirs(img_dir, exist_ok=True)
    mapping = {}
    for case_key in ("case1a", "case1b"):
        imgs = {}
        for img_key in ("IMG_Main", "IMG_Graph1", "IMG_Graph2"):
            fname = f"{case_key}_{img_key}.png"
            full = os.path.join(img_dir, fname)
            open(full, "wb").close()  # create empty file
            imgs[img_key] = full
        mapping[case_key] = imgs

    with open(path, "w", encoding="utf-8") as f:
        json.dump(mapping, f)
    return mapping


# ---------------------------------------------------------------------------
# ImageMappingValidator
# ---------------------------------------------------------------------------
class TestImageMappingValidator:
    def test_load_valid(self, tmp_path):
        img_dir = str(tmp_path / "imgs")
        mapping_path = str(tmp_path / "mapping.json")
        _make_mapping_json(mapping_path, img_dir)

        validator = ImageMappingValidator(base_dir=str(tmp_path))
        result = validator.load(mapping_path)
        assert "case1a" in result
        assert "IMG_Main" in result["case1a"]

    def test_missing_image_raises(self, tmp_path):
        bad_mapping = {"case1a": {"IMG_Main": "/no/such/image.png"}}
        mapping_path = str(tmp_path / "mapping.json")
        with open(mapping_path, "w") as f:
            json.dump(bad_mapping, f)

        validator = ImageMappingValidator()
        with pytest.raises(ValidationError, match="画像ファイルが見つかりません"):
            validator.load(mapping_path)

    def test_file_not_found(self):
        validator = ImageMappingValidator()
        with pytest.raises(ValidationError, match="not found"):
            validator.load("/no/such/mapping.json")


# ---------------------------------------------------------------------------
# DatasetBuilder
# ---------------------------------------------------------------------------
class TestDatasetBuilder:
    def test_build_structure(self, tmp_path):
        db_path = str(tmp_path / "work" / "database.db")
        _make_db(db_path)

        img_dir = str(tmp_path / "imgs")
        mapping = {}
        for case_key in ("case1a", "case1b"):
            imgs = {}
            for img_key in ("IMG_Main", "IMG_Graph1", "IMG_Graph2"):
                fname = f"{case_key}_{img_key}.png"
                full = os.path.join(img_dir, fname)
                os.makedirs(img_dir, exist_ok=True)
                open(full, "wb").close()
                imgs[img_key] = full
            mapping[case_key] = imgs

        builder = DatasetBuilder()
        dataset = builder.build(
            db_path=db_path,
            image_mapping=mapping,
            template_pptx_path="/fake/template.pptx",
        )

        assert dataset["template_pptx_path"].endswith("template.pptx")
        assert len(dataset["slides"]) == 2
        slide = dataset["slides"][0]
        assert slide["case_id"] == "1a"
        assert "TXT_ConfigFile" in slide["text_replacements"]
        assert "IMG_Main" in slide["image_replacements"]


# ---------------------------------------------------------------------------
# Phase2Pipeline (integration)
# ---------------------------------------------------------------------------
class TestPhase2Pipeline:
    def test_full_run(self, tmp_path):
        import yaml

        # Build config
        cfg = {
            "paths": {
                "db_path": str(tmp_path / "work" / "database.db"),
                "input_dir": str(tmp_path / "input"),
                "work_dir": str(tmp_path / "work"),
                "output_dir": str(tmp_path / "output"),
            },
            "sqlite": {"timeout": 30, "check_same_thread": False},
            "pptx": {"required_shapes": REQUIRED_SHAPES},
            "logging": {"level": "WARNING", "format": "%(message)s"},
        }
        cfg_path = str(tmp_path / "settings.yaml")
        with open(cfg_path, "w") as f:
            yaml.dump(cfg, f)

        # Build artefacts
        _make_db(str(tmp_path / "work" / "database.db"))

        pptx_path = str(tmp_path / "template.pptx")
        _make_pptx(pptx_path, REQUIRED_SHAPES)

        img_dir = str(tmp_path / "input" / "sample_data")
        mapping_path = str(tmp_path / "work" / "mapping.json")
        _make_mapping_json(mapping_path, img_dir)

        output_path = str(tmp_path / "output" / "ppt_dataset.json")

        pipeline = Phase2Pipeline(config_path=cfg_path)
        result = pipeline.run(
            image_mapping_path=mapping_path,
            template_pptx_path=pptx_path,
            output_json_path=output_path,
        )

        assert result["status"] == "success", result["message"]
        assert os.path.isfile(result["json_path"])

        with open(result["json_path"], encoding="utf-8") as f:
            data = json.load(f)
        assert len(data["slides"]) == 2
