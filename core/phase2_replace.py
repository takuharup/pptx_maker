"""Phase 2: SQLite + JSON image mapping → VBA execution dataset (JSON).

Classes
-------
ImageMappingValidator  – Validate image_mapping.json and resolve absolute paths.
TemplateValidator      – Check that required shapes exist in the PPTX template.
DatasetBuilder         – Build the ppt_dataset.json structure.
Phase2Pipeline         – Top-level orchestrator.
"""
from __future__ import annotations

import json
import logging
import os
import sqlite3
from typing import Any

from pptx import Presentation

from core.config_manager import ConfigManager
from core.validator import ValidationError, require_file

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# ImageMappingValidator
# ---------------------------------------------------------------------------
class ImageMappingValidator:
    """Load image_mapping.json and validate / absolutise all image paths."""

    def __init__(self, base_dir: str | None = None) -> None:
        # Images are resolved relative to this directory when the path is
        # relative.  Defaults to CWD.
        self._base_dir = base_dir or os.getcwd()

    def load(self, mapping_path: str) -> dict[str, dict[str, str]]:
        abs_mapping = require_file(mapping_path, "image_mapping.json")
        with open(abs_mapping, encoding="utf-8") as f:
            raw: dict[str, Any] = json.load(f)

        if not isinstance(raw, dict):
            raise ValidationError("image_mapping.json: トップレベルはオブジェクトである必要があります。")

        resolved: dict[str, dict[str, str]] = {}
        missing: list[str] = []

        for case_key, images in raw.items():
            if not isinstance(images, dict):
                raise ValidationError(
                    f"image_mapping.json: '{case_key}' の値はオブジェクトである必要があります。"
                )
            resolved_images: dict[str, str] = {}
            for img_key, img_path in images.items():
                abs_img = (
                    img_path
                    if os.path.isabs(img_path)
                    else os.path.join(self._base_dir, img_path)
                )
                abs_img = os.path.normpath(abs_img)
                if not os.path.isfile(abs_img):
                    missing.append(f"{case_key}/{img_key}: {abs_img}")
                resolved_images[img_key] = abs_img
            resolved[case_key] = resolved_images

        if missing:
            raise ValidationError(
                "以下の画像ファイルが見つかりません:\n" + "\n".join(missing)
            )

        return resolved


# ---------------------------------------------------------------------------
# TemplateValidator
# ---------------------------------------------------------------------------
class TemplateValidator:
    """Verify that a PPTX template contains all required shape names."""

    def validate(self, template_path: str, required_shapes: list[str]) -> None:
        abs_path = require_file(template_path, "Template PPTX")
        prs = Presentation(abs_path)

        found: set[str] = set()
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.name:
                    found.add(shape.name)

        missing = [s for s in required_shapes if s not in found]
        if missing:
            raise ValidationError(
                f"テンプレートに必須シェイプが存在しません: {missing}\n"
                f"検出されたシェイプ: {sorted(found)}"
            )
        logger.info("Template validation OK. Shapes found: %s", sorted(found))


# ---------------------------------------------------------------------------
# DatasetBuilder
# ---------------------------------------------------------------------------
class DatasetBuilder:
    """Build the ppt_dataset.json from DB rows and resolved image paths."""

    def build(
        self,
        db_path: str,
        image_mapping: dict[str, dict[str, str]],
        template_pptx_path: str,
        sqlite_timeout: int = 30,
    ) -> dict:
        cases = self._load_cases(db_path, sqlite_timeout)
        slides = []
        for case in cases:
            case_id = case["case_id"]
            # The JSON key can be either "case1a" or "1a" – try both.
            mapping_entry = image_mapping.get(f"case{case_id}") or image_mapping.get(case_id) or {}

            slide = {
                "case_id": case_id,
                "output_filename": f"output_case{case_id}.pptx",
                "text_replacements": self._build_text_replacements(case),
                "image_replacements": {
                    k: v for k, v in mapping_entry.items()
                    if k.startswith("IMG_")
                },
            }
            slides.append(slide)

        dataset = {
            "template_pptx_path": os.path.abspath(template_pptx_path),
            "slides": slides,
        }
        logger.info("Built dataset with %d slides.", len(slides))
        return dataset

    @staticmethod
    def _load_cases(db_path: str, timeout: int) -> list[dict]:
        abs_db = require_file(db_path, "SQLite DB")
        con = sqlite3.connect(abs_db, timeout=timeout)
        con.row_factory = sqlite3.Row
        try:
            rows = con.execute(
                "SELECT case_id, config_file, conditions, max_height, "
                "dispersion_height, gas_holdup FROM cases ORDER BY case_id"
            ).fetchall()
            return [dict(r) for r in rows]
        finally:
            con.close()

    @staticmethod
    def _build_text_replacements(case: dict) -> dict[str, str]:
        def fmt_float(val: float | None, decimals: int = 3) -> str:
            return f"{val:.{decimals}f}" if val is not None else "N/A"

        return {
            "TXT_ConfigFile": f"設定ファイル名：{case['config_file']}",
            "TXT_Conditions": case["conditions"] or "",
            "TXT_Dispersion": (
                f"時刻20 sec における90%分散高さ：{fmt_float(case['dispersion_height'])} m"
            ),
            "TXT_GasHU": (
                f"時刻20 sec におけるガスホールドアップ：{fmt_float(case['gas_holdup'])} %"
            ),
        }


# ---------------------------------------------------------------------------
# Phase2Pipeline
# ---------------------------------------------------------------------------
class Phase2Pipeline:
    """Orchestrate Phase 2: SQLite + JSON image mapping → ppt_dataset.json."""

    def __init__(self, config_path: str = "config/settings.yaml") -> None:
        self._cfg = ConfigManager(config_path)
        self._img_validator = ImageMappingValidator(
            base_dir=os.path.abspath(self._cfg.input_dir)
        )
        self._tpl_validator = TemplateValidator()
        self._builder = DatasetBuilder()

    def run(
        self,
        image_mapping_path: str,
        template_pptx_path: str,
        output_json_path: str,
    ) -> dict:
        """
        Execute Phase 2.

        Returns
        -------
        dict with keys: status, message, json_path
        """
        logger.info("=== Phase 2 start ===")
        try:
            image_mapping = self._img_validator.load(image_mapping_path)
            self._tpl_validator.validate(template_pptx_path, self._cfg.required_shapes)
            dataset = self._builder.build(
                db_path=self._cfg.db_path,
                image_mapping=image_mapping,
                template_pptx_path=template_pptx_path,
                sqlite_timeout=self._cfg.sqlite_timeout,
            )

            abs_output = os.path.abspath(output_json_path)
            os.makedirs(os.path.dirname(abs_output), exist_ok=True)
            with open(abs_output, "w", encoding="utf-8") as f:
                json.dump(dataset, f, ensure_ascii=False, indent=2)

            msg = f"Phase 2 完了: {abs_output} を生成しました。"
            logger.info(msg)
            return {"status": "success", "message": msg, "json_path": abs_output}

        except ValidationError as e:
            logger.error("Validation error: %s", e)
            return {"status": "error", "message": f"バリデーションエラー: {e}"}
        except Exception as e:
            logger.exception("Unexpected error in Phase 2")
            return {"status": "error", "message": f"予期しないエラー: {e}"}
